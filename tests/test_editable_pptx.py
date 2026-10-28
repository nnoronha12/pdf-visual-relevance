from __future__ import annotations

import base64
import io

from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from app.editable_pptx import embeddable_image
from app.main import app


def png_data_url(colour: tuple[int, int, int], size: tuple[int, int] = (40, 20)) -> str:
    buffer = io.BytesIO()
    Image.new("RGB", size, colour).save(buffer, format="PNG")
    return "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode()


LOGO = png_data_url((200, 30, 30))

HTML = f"""<!doctype html>
<html lang="en">
<head>
  <style>
    * {{ box-sizing: border-box; }}
    body {{ margin: 0; background: #ddd; }}
    .slide {{ position: relative; width: 1920px; height: 1080px; overflow: hidden; background: #fff; font-family: Arial, sans-serif; color: #404042; }}
    .title {{ position: absolute; left: 115px; top: 100px; width: 1400px; margin: 0; font-size: 60px; line-height: 1.1; font-weight: 400; }}
    .eyebrow {{ position: absolute; left: 115px; top: 60px; margin: 0; font-size: 18px; font-weight: 700; letter-spacing: .16em; text-transform: uppercase; color: #9A7430; }}
    .card {{ position: absolute; left: 115px; top: 400px; width: 500px; height: 260px; padding: 30px; background: #F2F3F3; border-radius: 12px; }}
    .rule {{ position: absolute; left: 115px; top: 300px; width: 88px; height: 4px; background: #B68A35; }}
    .row {{ position: absolute; left: 700px; top: 400px; width: 900px; padding: 20px 0; border-top: 1px solid #D6D9D9; font-size: 28px; }}
    .list {{ position: absolute; left: 700px; top: 600px; width: 900px; counter-reset: it; }}
    .item {{ display: grid; grid-template-columns: 80px 1fr; font-size: 28px; counter-increment: it; }}
    .item:before {{ content: counter(it, decimal-leading-zero); color: #B68A35; }}
    .logo {{ position: absolute; right: 115px; top: 60px; width: 160px; height: 80px; object-fit: contain; }}
    .scrimmed {{ position: absolute; left: 115px; top: 750px; width: 400px; height: 200px; background-image: url({LOGO}); background-size: cover; }}
    .scrimmed::after {{ content: ""; position: absolute; inset: 0; background: linear-gradient(90deg, rgba(0,0,0,.8), transparent); }}
    .slide[dir=rtl] {{ direction: rtl; }}
  </style>
</head>
<body>
  <section class="slide" data-layout="cover">
    <p class="eyebrow" data-slot="eyebrow">Ministry presentation</p>
    <h1 class="title" data-slot="title">Budget outlook <b>2027</b></h1>
    <div class="rule"></div>
    <div class="card"><span data-slot="stat">AED 71.5B</span></div>
    <div class="row">First row text</div>
    <div class="list"><div class="item">Alpha</div><div class="item">Beta</div></div>
    <img class="logo" src="{LOGO}" alt="Ministry logo">
    <div class="scrimmed"></div>
  </section>
  <section class="slide" dir="rtl" lang="ar">
    <h1 class="title" data-slot="title">آفاق الميزانية</h1>
  </section>
</body>
</html>"""


def payload(**extra: object) -> dict[str, object]:
    return {
        "html": HTML,
        "format": "pptx",
        "filename": "editable-test.pptx",
        "deck_id": "97000000-0000-4000-8000-000000000002",
        "revision": 1,
        **extra,
    }


def texts_of(slide) -> dict[str, object]:
    return {shape.text_frame.text: shape for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text.strip()}


def test_pptx_is_editable_by_default() -> None:
    with TestClient(app) as client:
        response = client.post("/convert", json=payload())
    assert response.status_code == 200
    assert response.headers["x-pptx-mode"] == "editable"
    presentation = Presentation(io.BytesIO(response.content))
    assert len(presentation.slides) == 2
    assert round(presentation.slide_width / 914400, 2) == 13.33
    assert round(presentation.slide_height / 914400, 2) == 7.5

    first, second = presentation.slides
    texts = texts_of(first)
    # Text is real, editable text, with the browser's casing and inline runs.
    assert "MINISTRY PRESENTATION" in texts
    assert "Budget outlook 2027" in texts
    title = texts["Budget outlook 2027"]
    assert title.name == "title"
    runs = title.text_frame.paragraphs[0].runs
    assert [r.text for r in runs] == ["Budget outlook ", "2027"]
    assert runs[1].font.bold is True
    assert runs[0].font.size.pt == 30  # 60 px on a 1920 px slide
    # CSS counters become editable numbers.
    assert "01" in texts and "02" in texts
    assert "Alpha" in texts and "AED 71.5B" in texts

    # The background is a picture; the plain boxes and rules are native shapes.
    assert first.shapes[0].name == "Background"
    autoshapes = [s for s in first.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not s.text_frame.text.strip()]
    fills = {s.fill.fore_color.rgb.__str__() for s in autoshapes if s.fill.type is not None}
    assert {"B68A35", "F2F3F3", "D6D9D9"} <= fills
    pictures = [s for s in first.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE and s.name != "Background"]
    # The logo is a picture of its own; the photo under the gradient scrim stays in the background.
    assert len(pictures) == 1
    assert pictures[0]._element.nvPicPr.cNvPr.get("descr") == "Ministry logo"

    arabic = texts_of(second)["آفاق الميزانية"]
    ppr = arabic.text_frame.paragraphs[0]._p.find(qn("a:pPr"))
    assert ppr.get("rtl") == "1"
    assert ppr.get("algn") == "r"
    assert arabic.text_frame.paragraphs[0].runs[0]._r.find(qn("a:rPr")).get("lang") == "ar-AE"


def test_image_mode_keeps_one_picture_per_slide() -> None:
    with TestClient(app) as client:
        response = client.post("/convert", json=payload(pptx_mode="image"))
    assert response.status_code == 200
    assert response.headers["x-pptx-mode"] == "image"
    presentation = Presentation(io.BytesIO(response.content))
    for slide in presentation.slides:
        assert [s.shape_type for s in slide.shapes] == [MSO_SHAPE_TYPE.PICTURE]


def test_html_without_slides_falls_back_to_pictures() -> None:
    html = "<!doctype html><html><body><h1>No slide markup</h1></body></html>"
    with TestClient(app) as client:
        response = client.post("/convert", json={**payload(), "html": html})
    assert response.status_code == 200
    assert response.headers["x-pptx-mode"] == "image"
    assert len(Presentation(io.BytesIO(response.content)).slides) >= 1


def test_rejects_unknown_pptx_mode() -> None:
    response = TestClient(app).post("/convert", json=payload(pptx_mode="vector"))
    assert response.status_code == 422


def test_embeddable_image_converts_webp() -> None:
    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), (1, 2, 3)).save(buffer, format="WEBP")
    data, image_format = embeddable_image(buffer.getvalue())
    assert image_format == "PNG" and data.startswith(b"\x89PNG")
    assert embeddable_image(b"<svg xmlns='http://www.w3.org/2000/svg'/>") is None
