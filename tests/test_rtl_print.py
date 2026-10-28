from __future__ import annotations

import io

import fitz
from fastapi.testclient import TestClient
from pptx import Presentation

from app.main import app

# The deck shell's layout: 1920x1080 slides inside narrower frames, scaled from
# the top-left corner for print. In an Arabic (dir="rtl") deck the slides used
# to land off the page and spill over two pages each.
HTML = """<!doctype html>
<html lang="ar" dir="rtl">
<head>
  <style>
    body { margin: 0; }
    .deck { display: flex; flex-direction: column; gap: 28px; align-items: center; padding: 32px 16px; }
    .frame { width: 1100px; overflow: hidden; }
    .frame > .slide { transform-origin: top left; }
    .slide { position: relative; width: 1920px; height: 1080px; background: #404042; color: #fff; font: 80px Arial, sans-serif; }
    .slide h1 { position: absolute; top: 400px; right: 115px; margin: 0; font-size: 80px; }
    @media print {
      @page { size: 13.333in 7.5in; margin: 0; }
      .deck { display: block; padding: 0; }
      .frame { width: 13.333in !important; height: 7.5in !important; break-after: page; }
      .frame > .slide { transform: scale(calc(13.333 * 96 / 1920)) !important; }
    }
  </style>
</head>
<body>
  <div class="deck">
    <div class="frame"><section class="slide" lang="ar" dir="rtl"><h1>الشريحة 1</h1></section></div>
    <div class="frame"><section class="slide" lang="ar" dir="rtl"><h1>الشريحة 2</h1></section></div>
    <div class="frame"><section class="slide" lang="ar" dir="rtl"><h1>الشريحة 3</h1></section></div>
  </div>
</body>
</html>"""


def convert(output_format: str, **extra: object):
    with TestClient(app) as client:
        return client.post("/convert", json={
            "html": HTML,
            "format": output_format,
            "filename": f"arabic.{output_format}",
            "deck_id": "97000000-0000-4000-8000-000000000003",
            "revision": 0,
            **extra,
        })


def test_arabic_pdf_keeps_one_whole_slide_per_page() -> None:
    response = convert("pdf")
    assert response.status_code == 200
    document = fitz.open(stream=response.content, filetype="pdf")
    try:
        assert document.page_count == 3
        for page, expected in zip(document, ["1", "2", "3"]):
            words = page.get_text("words")
            assert any(w[4] == expected for w in words), page.get_text()
            # The heading sits near the right edge, inside the page.
            x1 = max(w[2] for w in words)
            assert page.rect.width * 0.8 < x1 <= page.rect.width
            # The slide background fills the page instead of a sliver of it.
            pixel = page.get_pixmap(dpi=10).pixel(2, 2)
            assert all(abs(a - b) <= 3 for a, b in zip(pixel, (64, 64, 66))), pixel
    finally:
        document.close()


def test_arabic_pptx_in_image_mode_has_one_slide_per_page() -> None:
    response = convert("pptx", pptx_mode="image")
    assert response.status_code == 200
    assert len(Presentation(io.BytesIO(response.content)).slides) == 3


def test_arabic_pptx_editable_text_stays_on_the_slide() -> None:
    response = convert("pptx")
    assert response.status_code == 200
    assert response.headers["x-pptx-mode"] == "editable"
    presentation = Presentation(io.BytesIO(response.content))
    assert len(presentation.slides) == 3
    for slide in presentation.slides:
        boxes = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        assert len(boxes) == 1
        box = boxes[0]
        assert 0 < box.left and box.left + box.width <= presentation.slide_width
