"""HTML slides to an editable PPTX.

Each slide is rendered in Chromium and measured by ``editable_pptx.js``. Text
becomes PowerPoint text boxes, plain boxes and rules become native shapes, and
photos and logos become pictures. Anything that cannot be rebuilt natively
(gradients, patterns, rotated decoration) stays in a background picture that is
captured after the rebuilt parts are hidden, so nothing is drawn twice.
"""

from __future__ import annotations

import asyncio
import base64
import io
import logging
import re
from pathlib import Path
from typing import Any
from urllib.parse import unquote_to_bytes

from lxml import etree
from PIL import Image
from playwright.async_api import Browser, Error as PlaywrightError, Page, TimeoutError as PlaywrightTimeoutError
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Emu

from .html_conversion import RENDER_TIMEOUT_MS, HTMLConversionError, load_html

logger = logging.getLogger(__name__)

PAGE_SCRIPT = (Path(__file__).with_name("editable_pptx.js")).read_text(encoding="utf-8")

SLIDE_WIDTH_EMU = 12_192_000  # 13.333 in, PowerPoint's 16:9 width
EMU_PER_PT = 12_700
MAX_IMAGE_BYTES = 25 * 1024 * 1024
BACKGROUND_JPEG_THRESHOLD = 1_200_000
EMBEDDABLE_FORMATS = {"PNG", "JPEG", "GIF", "BMP", "TIFF"}

ALIGN = {"left": "l", "right": "r", "center": "ctr", "justify": "just"}


class NoSlidesFound(Exception):
    """The HTML has no `.slide` elements to rebuild."""


# --------------------------------------------------------------------------- browser


async def _collect_images(page: Page, store: dict[str, bytes], tasks: list[asyncio.Task]) -> None:
    async def keep(response) -> None:
        try:
            if response.request.resource_type != "image" or not response.ok:
                return
            body = await response.body()
            if 0 < len(body) <= MAX_IMAGE_BYTES:
                store[response.url] = body
        except Exception:  # a failed image only means it stays in the background picture
            pass

    page.on("response", lambda response: tasks.append(asyncio.ensure_future(keep(response))))


def _decode_data_url(url: str) -> bytes | None:
    match = re.match(r"data:([^,;]*)((?:;[^,;]*)*),(.*)$", url, re.S)
    if not match:
        return None
    payload = match.group(3)
    if ";base64" in match.group(2):
        try:
            return base64.b64decode(payload)
        except Exception:
            return None
    return unquote_to_bytes(payload)


def embeddable_image(data: bytes) -> tuple[bytes, str] | None:
    """Returns image bytes PowerPoint can embed (converting e.g. WebP to PNG), or None."""
    try:
        with Image.open(io.BytesIO(data)) as image:
            image_format = (image.format or "").upper()
            if image_format in EMBEDDABLE_FORMATS:
                return data, image_format
            image.load()
            out = io.BytesIO()
            image.save(out, format="PNG")
            return out.getvalue(), "PNG"
    except Exception:
        return None


def _background_bytes(png: bytes) -> bytes:
    if len(png) <= BACKGROUND_JPEG_THRESHOLD:
        return png
    with Image.open(io.BytesIO(png)) as image:
        out = io.BytesIO()
        image.convert("RGB").save(out, format="JPEG", quality=90, optimize=True)
        return out.getvalue()


async def extract_slides(browser: Browser, html: str) -> tuple[list[dict[str, Any]], dict[str, bytes]]:
    """Measures every slide and captures its background. Returns (slides, images by URL)."""
    context = await browser.new_context(viewport={"width": 1920, "height": 1080}, device_scale_factor=1)
    page = await context.new_page()
    page.set_default_timeout(RENDER_TIMEOUT_MS)
    responses: dict[str, bytes] = {}
    tasks: list[asyncio.Task] = []
    await _collect_images(page, responses, tasks)
    try:
        await load_html(page, html)
        await page.evaluate(PAGE_SCRIPT)
        sizes = await page.evaluate("() => window.__pptx.prepare()")
        if not sizes or any(s["w"] <= 0 or s["h"] <= 0 for s in sizes):
            raise NoSlidesFound()
        width = max(int(s["w"]) for s in sizes)
        height = max(int(s["h"]) for s in sizes)
        if (width, height) != (1920, 1080):
            await page.set_viewport_size({"width": width, "height": height})

        if tasks:
            await asyncio.gather(*tasks, return_exceptions=True)
        images: dict[str, bytes] = {}
        for url in await page.evaluate("() => window.__pptx.sources()"):
            data = responses.get(url) or (_decode_data_url(url) if url.startswith("data:") else None)
            converted = embeddable_image(data) if data else None
            if converted:
                images[url] = converted[0]

        slides: list[dict[str, Any]] = []
        for index in range(len(sizes)):
            size = await page.evaluate("(i) => window.__pptx.focus(i)", index)
            info = await page.evaluate("([i, usable]) => window.__pptx.extract(i, usable)", [index, list(images)])
            shot = await page.screenshot(clip={"x": 0, "y": 0, "width": size["w"], "height": size["h"]}, type="png", animations="disabled")
            info["background"] = _background_bytes(shot)
            slides.append(info)
            await page.evaluate("(i) => window.__pptx.unfocus(i)", index)
        return slides, images
    except (PlaywrightError, PlaywrightTimeoutError) as exc:
        raise HTMLConversionError(f"Chromium could not render the HTML: {exc}") from exc
    finally:
        await context.close()


# --------------------------------------------------------------------------- PPTX


def _srgb(hex_value: str, alpha: float) -> etree._Element:
    clr = etree.SubElement(etree.Element(qn("a:solidFill")), qn("a:srgbClr"), val=hex_value)
    if alpha < 0.999:
        etree.SubElement(clr, qn("a:alpha"), val=str(max(0, min(100_000, round(alpha * 100_000)))))
    return clr.getparent()


def _set_fill(shape, fill: dict | None) -> None:
    sp_pr = shape._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:gradFill", "a:pattFill", "a:blipFill"):
        for old in sp_pr.findall(qn(tag)):
            sp_pr.remove(old)
    node = _srgb(fill["hex"], fill["alpha"]) if fill else etree.Element(qn("a:noFill"))
    # Fill goes right after the geometry element.
    geom = sp_pr.find(qn("a:prstGeom"))
    geom.addnext(node)


def _set_line(shape, line: dict | None, emu_per_px: float) -> None:
    sp_pr = shape._element.spPr
    for old in sp_pr.findall(qn("a:ln")):
        sp_pr.remove(old)
    ln = etree.SubElement(sp_pr, qn("a:ln"))
    if line:
        ln.set("w", str(max(1, round(line["width"] * emu_per_px))))
        ln.append(_srgb(line["hex"], line["alpha"]))
    else:
        etree.SubElement(ln, qn("a:noFill"))


def _drop_style(shape) -> None:
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _rect_emu(obj: dict, emu_per_px: float) -> tuple[Emu, Emu, Emu, Emu]:
    return (
        Emu(round(obj["x"] * emu_per_px)),
        Emu(round(obj["y"] * emu_per_px)),
        Emu(max(1, round(obj["w"] * emu_per_px))),
        Emu(max(1, round(obj["h"] * emu_per_px))),
    )


def _add_shape(slide, obj: dict, emu_per_px: float, number: int) -> None:
    kind = {"rect": MSO_SHAPE.RECTANGLE, "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL}[obj["kind"]]
    shape = slide.shapes.add_shape(kind, *_rect_emu(obj, emu_per_px))
    shape.name = f"Shape {number}"
    _drop_style(shape)
    if obj["kind"] == "roundRect":
        shortest = max(1e-6, min(obj["w"], obj["h"]))
        shape.adjustments[0] = max(0.0, min(0.5, obj.get("radius", 0) / shortest))
    _set_fill(shape, obj.get("fill"))
    _set_line(shape, obj.get("line"), emu_per_px)


def _add_picture(slide, obj: dict, data: bytes, emu_per_px: float, number: int) -> None:
    picture = slide.shapes.add_picture(io.BytesIO(data), *_rect_emu(obj, emu_per_px))
    crop = obj.get("crop") or {}
    picture.crop_left = crop.get("l", 0)
    picture.crop_right = crop.get("r", 0)
    picture.crop_top = crop.get("t", 0)
    picture.crop_bottom = crop.get("b", 0)
    picture.name = f"Picture {number}"
    if obj.get("alt"):
        picture._element.nvPicPr.cNvPr.set("descr", obj["alt"][:1000])


def _run_properties(run: dict, pt_per_px: float) -> etree._Element:
    rpr = etree.Element(qn("a:rPr"))
    rpr.set("lang", run.get("lang") or "en-US")
    rpr.set("sz", str(max(100, round(run["size"] * pt_per_px * 100))))
    rpr.set("b", "1" if run.get("bold") else "0")
    rpr.set("i", "1" if run.get("italic") else "0")
    if run.get("underline"):
        rpr.set("u", "sng")
    if run.get("letterSpacing"):
        rpr.set("spc", str(round(run["letterSpacing"] * pt_per_px * 100)))
    rpr.set("dirty", "0")
    rpr.append(_srgb(run.get("color") or "000000", float(run.get("alpha", 1))))
    font = run.get("font") or "Arial"
    etree.SubElement(rpr, qn("a:latin"), typeface=font)
    etree.SubElement(rpr, qn("a:cs"), typeface=font)
    return rpr


def _add_text(slide, box: dict, emu_per_px: float, pt_per_px: float) -> None:
    shape = slide.shapes.add_textbox(*_rect_emu(box, emu_per_px))
    shape.name = str(box.get("name") or "Text")[:60]
    frame = shape.text_frame
    frame.word_wrap = bool(box.get("wrap"))
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    body_pr = frame._txBody.bodyPr
    body_pr.set("anchor", "t")
    if box.get("rtl"):
        body_pr.set("rtlCol", "0")

    line_pts = round(float(box.get("lineHeight") or 0) * pt_per_px * 100)
    for index, runs in enumerate(box["paragraphs"]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p = paragraph._p
        ppr = p.get_or_add_pPr()
        ppr.set("algn", ALIGN.get(box.get("align"), "l"))
        if box.get("rtl"):
            ppr.set("rtl", "1")
        if line_pts > 0:
            ln_spc = etree.SubElement(ppr, qn("a:lnSpc"))
            etree.SubElement(ln_spc, qn("a:spcPts"), val=str(line_pts))
        for run in runs:
            r = paragraph.add_run()
            r.text = run["text"]
            old = r._r.find(qn("a:rPr"))
            if old is not None:
                r._r.remove(old)
            r._r.insert(0, _run_properties(run, pt_per_px))
        if not runs:
            end = p.find(qn("a:endParaRPr"))
            if end is None:
                etree.SubElement(p, qn("a:endParaRPr"), lang="en-US", dirty="0")


def build_pptx(slides: list[dict[str, Any]], images: dict[str, bytes], title: str = "") -> bytes:
    if not slides:
        raise HTMLConversionError("The HTML rendered without any slides")
    first = slides[0]
    emu_per_px = SLIDE_WIDTH_EMU / first["w"]
    pt_per_px = emu_per_px / EMU_PER_PT

    presentation = Presentation()
    presentation.slide_width = Emu(SLIDE_WIDTH_EMU)
    presentation.slide_height = Emu(round(first["h"] * emu_per_px))
    if title:
        presentation.core_properties.title = title[:255]
    blank = presentation.slide_layouts[6]

    for info in slides:
        slide = presentation.slides.add_slide(blank)
        background = slide.shapes.add_picture(io.BytesIO(info["background"]), 0, 0, presentation.slide_width, presentation.slide_height)
        background.name = "Background"
        number = 0
        for obj in info.get("objects", []):
            number += 1
            if obj["kind"] == "picture":
                data = images.get(obj["src"])
                if data:
                    _add_picture(slide, obj, data, emu_per_px, number)
            else:
                _add_shape(slide, obj, emu_per_px, number)
        for box in info.get("texts", []):
            _add_text(slide, box, emu_per_px, pt_per_px)

    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


async def render_html_to_editable_pptx(browser: Browser, html: str, title: str = "") -> bytes:
    slides, images = await extract_slides(browser, html)
    return build_pptx(slides, images, title)
